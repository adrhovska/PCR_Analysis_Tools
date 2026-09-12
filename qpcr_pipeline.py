#!/usr/bin/env python3
"""
qPCR analysis tool which employs the updated SD-based replicate removal method
(optionally skippable via --skip-outlier-removal, see step 2 below).
Allows for exploratory imputation of failed (sample, target) groups using several strategies.
Includes statistical testing (parametric or non-parametric) and plotting of results.

WORKFLOW
1. Reads one or more QuantStudio qPCR export files (one file per plate). Each
   file must contain a "Results" sheet with the standard per-well columns
   (Well, Well Position, Sample, Target, Cq, ...).
2. Groups technical replicates by (plate, sample, target) and applies an
   SD-based outlier removal rule:
     - compute SD of the replicate group
     - if SD <= set threshold -> pass, keep all replicates
     - if SD > threshold  -> remove the replicate farthest from the group
       mean, recompute SD, repeat until below threshold
     - if only 2 replicates remain and SD is still > threshold -> the
       whole (sample, target) group FAILS QC (but allow exploration later)
   The SD threshold is a user parsed CLI argument (--sd-threshold).
   Pass --skip-outlier-removal to bypass this step entirely: every valid
   (non-NaN) replicate is kept and simply averaged, and a group only FAILS
   if every replicate had no Cq at all (no amplification). --sd-threshold
   is ignored in that mode.
   Biological condition + replicate id are resolved per Sample using, in
   priority order: an explicit --sample-map CSV, then DA3's own 'Biogroup'
   field (Replicate Group Result sheet) if it was actually populated during
   plate setup, then a regex on the Sample text (--condition-regex, default:
   a trailing " <number>"). Sample names do NOT need to encode the replicate
   number themselves as long as DA3's Biogroup field was set, or you supply
   --sample-map.
3. In "exploratory" mode (--exploratory), failed (sample, target) groups
   can be salvaged instead of dropped, using one of several strategies
   (--fail-strategy): plate_mean, anchor_mean, best_subset, or drop.
   This is most useful for housekeeping genes, where losing a sample
   entirely because of one bad technical replicate is wasteful but
   exploratory imputation is reasonable/desired.
4. Computes ddCt per plate, referenced to a user-specified anchor SAMPLE
   (one specific well/biological replicate used as the calibrator, e.g.
   "Empty 1" or "Church wt"). Giving --anchor is what triggers this step
   (and step 5) at all -- omit it entirely to get QC + raw Cq outputs only.
   Two modes, depending on whether --housekeeping is also given:
     - WITH a housekeeping gene (e.g. GAPDH) -- the standard two-step method:
           dCt        = Cq(target) - Cq(housekeeping)
           ddCt       = dCt(sample) - dCt(anchor)
     - WITHOUT a housekeeping gene (--housekeeping omitted) -- ddCt is taken
       directly off each target's own Cq, with no reference-gene
       normalization step at all:
           dCt        = Cq(target)          (no housekeeping subtraction)
           ddCt       = dCt(sample) - dCt(anchor)
       Use this mode when there is no housekeeping/endogenous-control assay
       on the plate by design -- e.g. template input was already equalized
       another way (Qubit-based dilution to a fixed ng amount), so forcing a
       reference-gene normalization would just add noise rather than remove it.
   Either way:
       RQ         = 2^-ddCt
       log2FC     = -ddCt   (== log2(RQ))
5. Runs group-wise statistics on log2FC per target vs. a control condition
   (all biological replicates of the anchor's condition, not just the
   single anchor well): per-group Shapiro-Wilk normality + Levene's test
   for equal variances decide parametric vs. non-parametric, then:
     - parametric:      one-way ANOVA + Dunnett's post-hoc vs. control
     (can be changes but ANOVA used in previous analysis and this allows comparison)
     - non-parametric:  Kruskal-Wallis + Dunn's post-hoc vs. control
6. Writes result tables (technical-replicate level and
   sample/target-level with Cq, dCt, ddCt, RQ, log2FC) plus a QC log,
   and generates plots:
     - qc/            per-plate Cq-per-technical-replicate plots with
                       removed points marked (X), faceted by target
     - cq/             Cq per sample per target, raw and anchor-referenced
     - log2fc/         log2FC per condition per target with significance
                       brackets vs. the control condition -- one file per
                       target (log2FC_<target>.png) plus a single combined
                       figure with all targets side by side
                       (log2FC_all_targets.png)
     - rq/             the same comparison as log2fc/, drawn instead as a
                       bar chart on the linear RQ scale: mean bar (= 2^-mean
                       ddCt) + whiskers (+/- 1 SEM of ddCt) + individual
                       biological-replicate points + a dashed RQ=1 line,
                       with the same significance brackets -- one file per
                       target (RQ_<target>.png) plus a combined figure
                       (RQ_all_targets.png). Pass --rq-highlight-target-match
                       to color a condition's bar differently when the
                       target/gene name appears in the condition name (the
                       "targets this gene" vs. "does not target this gene"
                       convention for dCas9/CRISPRoff-style experiments).
   In qc/ and cq/, samples are ordered left-to-right by their physical
   position on the plate (first well occupied, row letter then column
   number) rather than alphabetically or by condition -- matches the
   plate layout, which makes visual interpretation easier.
   Plot titles are kept short ("QC -- {plate}", "{plate} -- Cq",
   "{target} vs {control}") so they fit at the default figure size; if a
   plate label or target name is still unusually long, the figure widens
   itself just enough to fit the title rather than clipping it.

USAGE (PARSING CLI)
    python qpcr_pipeline.py \\
        --input plate1.xlsx plate2.xlsx plate3.xlsx \\
        --plate-labels Plate1 Plate2 Plate3 \\
        --anchor "Empty 1" "E 1" "E 1" \\
        --housekeeping GAPDH \\
        --sd-threshold 0.3 \\
        --exploratory --fail-strategy plate_mean \\
        --outdir ./qpcr_results

    # Same run, but skip SD-based outlier removal entirely (average every
    # valid replicate as-is; --sd-threshold is ignored) and go straight to ddCt:
    python qpcr_pipeline.py \\
        --input plate1.xlsx plate2.xlsx plate3.xlsx \\
        --plate-labels Plate1 Plate2 Plate3 \\
        --anchor "Empty 1" "E 1" "E 1" \\
        --housekeeping GAPDH \\
        --skip-outlier-removal \\
        --outdir ./qpcr_results_no_qc

    # No housekeeping/endogenous-control assay at all (e.g. a Golden Gate
    # cutting-efficiency plate where template was Qubit-normalized): give
    # --anchor and simply omit --housekeeping -- ddCt is then computed
    # directly off each target's own Cq vs. the anchor sample:
    python qpcr_pipeline.py \\
        --input plate1.xlsx \\
        --anchor "Church wt" \\
        --outdir ./qpcr_results_no_housekeeping

Run `python qpcr_pipeline.py --help` for the full argument list.

DEPENDENCIES
    pip install pandas numpy openpyxl scipy matplotlib seaborn scikit-posthocs statsmodels
(scipy.stats.dunnett requires scipy >= 1.11)
"""

from __future__ import annotations

import argparse
import logging
import re
import sys
import warnings
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import seaborn as sns

from scipy import stats as sstats

try:
    import scikit_posthocs as sp
except ImportError:  # pragma: no cover
    sp = None

sns.set_theme(style="whitegrid", context="talk")

LOG = logging.getLogger("qpcr_pipeline")

STAR_THRESHOLDS = [(1e-4, "****"), (1e-3, "***"), (1e-2, "**"), (5e-2, "*")] 
def stars_for(p: Optional[float]) -> str:
    if p is None or np.isnan(p):
        return "n/a"
    for cutoff, label in STAR_THRESHOLDS:
        if p < cutoff:
            return label
    return "ns"


# 1. Reading QuantStudio "Results" sheets (has to be formatted beforehand)

REQUIRED_COLS = ["Well Position", "Sample", "Target", "Cq", "Omit", "Amp Status"]


def find_results_header_row(ws) -> int:
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row and row[0] == "Well":
            return i
    raise ValueError("Could not find the 'Well' header row in the Results sheet.")


def read_plate_results(path: Path, plate_label: str) -> pd.DataFrame:
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, data_only=True)
    if "Results" not in wb.sheetnames:
        raise ValueError(f"{path.name}: no 'Results' sheet found (sheets: {wb.sheetnames})")

    ws = wb["Results"]
    header_row = find_results_header_row(ws)
    headers = [c.value for c in ws[header_row]]

    records = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if row[0] is None:
            continue
        records.append(row)

    df = pd.DataFrame(records, columns=headers)

    missing = [c for c in REQUIRED_COLS if c not in df.columns]
    if missing:
        raise ValueError(f"{path.name}: Results sheet is missing expected columns: {missing}")

    df["Cq"] = pd.to_numeric(df["Cq"], errors="coerce")
    df["Sample"] = df["Sample"].astype(str).str.strip()
    df["Target"] = df["Target"].astype(str).str.strip()
    df["Well Position"] = df["Well Position"].astype(str).str.strip()
    df["Omit"] = df["Omit"].astype(str).str.upper().eq("TRUE")
    df["Plate"] = plate_label
    df["SourceFile"] = path.name

    keep_cols = ["Plate", "SourceFile", "Well Position", "Sample", "Target",
                 "Cq", "Omit", "Amp Status"]
    return df[keep_cols].reset_index(drop=True)


# 2. Condition / biological-replicate resolution (Sample does NOT have to be
#    named with a trailing replicate number -- see ConditionResolver below)

DEFAULT_CONDITION_REGEX = r"^(?P<condition>.*?)\s+(?P<rep>\d+)$"


def find_generic_header_row(ws, first_cell_options: set[str]) -> Optional[int]:
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row and row[0] in first_cell_options:
            return i
    return None


def read_biogroup_map(path: Path) -> dict:
    """Read DA3's own biological-replicate grouping, if it was actually used.

    QuantStudio Design & Analysis (DA3) exports a 'Replicate Group Result'
    sheet with a 'Biogroup' column, populated only if a "Biological Group"
    was assigned during plate setup in the DA3 app -- independent of the
    free-text 'Sample' name. If that field was never set, the column is
    entirely empty and this returns {}, so the caller falls back to
    --sample-map / --condition-regex.
    """
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, data_only=True)

    if "Replicate Group Result" not in wb.sheetnames:
        return {}
    ws = wb["Replicate Group Result"]
    header_row = find_generic_header_row(ws, {"Sample"})
    if header_row is None:
        return {}
    headers = [c.value for c in ws[header_row]]
    if "Sample" not in headers or "Biogroup" not in headers:
        return {}
    sample_idx = headers.index("Sample")
    biogroup_idx = headers.index("Biogroup")

    mapping = {}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if row[sample_idx] is None:
            continue
        bg = row[biogroup_idx]
        if bg is not None and str(bg).strip() != "":
            mapping[str(row[sample_idx]).strip()] = str(bg).strip()
    return mapping


def load_sample_map(path: Path) -> pd.DataFrame:
    """Explicit Sample -> Condition/BioRep override table (columns: Sample,
    Condition, BioRep, and an optional Plate column -- blank Plate matches
    any plate). Highest-priority resolution source; use this when Sample
    names carry no replicate info at all and DA3's Biogroup wasn't set.
    """
    df = pd.read_csv(path, dtype=str)
    required = {"Sample", "Condition", "BioRep"}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"--sample-map {path}: missing required column(s): {missing}")
    if "Plate" not in df.columns:
        df["Plate"] = np.nan
    return df


class ConditionResolver:
    """Resolves (plate, sample) -> (condition, bio_rep, source).

    Priority, highest first:
      1. an explicit --sample-map entry (per-plate, or wildcard across plates)
      2. DA3's own Biogroup field (per plate), when it was actually populated
      3. a regex applied to the Sample string (default or --condition-regex)
    """

    def __init__(self, sample_map: Optional[pd.DataFrame], biogroup_maps: dict,
                 condition_regex: str):
        self.sample_map = sample_map
        self.biogroup_maps = biogroup_maps  # plate -> {sample: biogroup}
        self.regex = re.compile(condition_regex)
        self.used_sources = set()

    def resolve(self, plate: str, sample: str) -> tuple[str, str, str]:
        if self.sample_map is not None:
            hit = self.sample_map[
                (self.sample_map["Sample"] == sample) &
                (self.sample_map["Plate"].isna() | (self.sample_map["Plate"] == plate))
            ]
            if not hit.empty:
                row = hit.iloc[0]
                self.used_sources.add("sample_map")
                return row["Condition"], row["BioRep"], "sample_map"

        bg_map = self.biogroup_maps.get(plate, {})
        if sample in bg_map:
            self.used_sources.add("biogroup")
            return bg_map[sample], sample, "biogroup"

        m = self.regex.match(sample)
        self.used_sources.add("regex")
        if m:
            gd = m.groupdict()
            return gd.get("condition", sample), gd.get("rep", "1"), "regex"
        return sample, "1", "regex"


# 3. SD-based outlier removal

@dataclass
class ReplicateResolution:
    plate: str
    sample: str
    target: str
    condition: str
    bio_rep: str
    kept_wells: list = field(default_factory=list)
    kept_cqs: list = field(default_factory=list)
    removed_wells: list = field(default_factory=list)
    removed_cqs: list = field(default_factory=list)
    removed_reasons: list = field(default_factory=list)
    n_start: int = 0
    final_sd: Optional[float] = None
    status: str = "PASS"          # PASS | PASS_AFTER_REMOVAL | PASS_NO_QC | FAILED
    fail_reason: Optional[str] = None
    cq_mean: Optional[float] = None
    imputed: bool = False
    imputed_note: Optional[str] = None


def resolve_replicate_group(plate, sample, target, condition, bio_rep,
                             wells, cqs, sd_threshold,
                             skip_outlier_removal: bool = False) -> ReplicateResolution:
    res = ReplicateResolution(plate=plate, sample=sample, target=target,
                               condition=condition, bio_rep=bio_rep)

    remaining = list(zip(wells, cqs))
    n_raw = len(remaining)

    # Drop replicates that never amplified (Cq is NaN) before SD logic starts, count as removed
    valid = [(w, c) for w, c in remaining if not pd.isna(c)]
    no_amp = [(w, c) for w, c in remaining if pd.isna(c)]
    for w, c in no_amp:
        res.removed_wells.append(w)
        res.removed_cqs.append(c)
        res.removed_reasons.append("no Cq / no amplification")

    remaining = valid
    res.n_start = n_raw

    if skip_outlier_removal:
        # Bypass the SD-based removal loop entirely: keep every valid
        # replicate and just average it. Only fail if nothing amplified.
        if not remaining:
            res.status = "FAILED"
            res.fail_reason = "no valid Cq values (no amplification in any replicate)"
            res.final_sd = np.nan
        else:
            cq_vals = np.array([c for _, c in remaining], dtype=float)
            res.final_sd = float(np.std(cq_vals, ddof=1)) if len(cq_vals) >= 2 else np.nan
            res.status = "PASS_NO_QC"
        res.kept_wells = [w for w, _ in remaining]
        res.kept_cqs = [c for _, c in remaining]
        if remaining:
            res.cq_mean = float(np.mean([c for _, c in remaining]))
        return res

    while True:
        if len(remaining) < 2:
            res.status = "FAILED"
            res.fail_reason = "fewer than 2 valid replicates available"
            res.final_sd = np.nan
            break

        cq_vals = np.array([c for _, c in remaining], dtype=float)
        sd = float(np.std(cq_vals, ddof=1))

        if sd <= sd_threshold:
            res.status = "PASS" if not res.removed_wells else "PASS_AFTER_REMOVAL"
            res.final_sd = sd
            break

        if len(remaining) == 2:
            res.status = "FAILED"
            res.fail_reason = (
                f"SD={sd:.3f} still above threshold ({sd_threshold}) "
                f"with only 2 replicates left"
            )
            res.final_sd = sd
            break

        # remove the replicate farthest from the current mean
        mean = float(np.mean(cq_vals))
        idx = int(np.argmax(np.abs(cq_vals - mean)))
        bad_well, bad_cq = remaining.pop(idx)
        res.removed_wells.append(bad_well)
        res.removed_cqs.append(bad_cq)
        res.removed_reasons.append(
            f"largest deviation from mean ({abs(bad_cq - mean):.3f} Cq) while SD={sd:.3f}"
        )

    res.kept_wells = [w for w, _ in remaining]
    res.kept_cqs = [c for _, c in remaining]
    if remaining:
        res.cq_mean = float(np.mean([c for _, c in remaining]))
    return res


def run_qc(raw: pd.DataFrame, sd_threshold: float,
           resolver: "ConditionResolver",
           skip_outlier_removal: bool = False) -> tuple[pd.DataFrame, pd.DataFrame]:
    """replicate_log: one row per technical replicate (kept/removed, with reason).
    group_summary: one row per (plate, sample, target) with QC outcome + Cq mean."""
    rep_rows = []
    group_rows = []

    grouped = raw.groupby(["Plate", "Sample", "Target"], sort=False)
    for (plate, sample, target), sub in grouped:
        condition, bio_rep, _source = resolver.resolve(plate, sample)
        res = resolve_replicate_group(
            plate, sample, target, condition, bio_rep,
            sub["Well Position"].tolist(), sub["Cq"].tolist(),
            sd_threshold, skip_outlier_removal=skip_outlier_removal,
        )

        for w, c in zip(res.kept_wells, res.kept_cqs):
            rep_rows.append(dict(Plate=plate, Sample=sample, Condition=condition,
                                  BioRep=bio_rep, Target=target, Well=w, Cq=c,
                                  Removed=False, Reason=None))
        for w, c, r in zip(res.removed_wells, res.removed_cqs, res.removed_reasons):
            rep_rows.append(dict(Plate=plate, Sample=sample, Condition=condition,
                                  BioRep=bio_rep, Target=target, Well=w, Cq=c,
                                  Removed=True, Reason=r))

        group_rows.append(dict(
            Plate=plate, Sample=sample, Condition=condition, BioRep=bio_rep,
            Target=target, n_start=res.n_start, n_kept=len(res.kept_wells),
            n_removed=len(res.removed_wells), QC_status=res.status,
            fail_reason=res.fail_reason, Cq_mean=res.cq_mean,
            Cq_SD=res.final_sd, Imputed=False, Imputed_note=None,
        ))

    replicate_log = pd.DataFrame(rep_rows)
    group_summary = pd.DataFrame(group_rows)
    return replicate_log, group_summary

# 4. Exploratory imputation of FAILED groups

def apply_fail_strategy(group_summary: pd.DataFrame, raw: pd.DataFrame,
                         exploratory: bool, fail_strategy: str,
                         anchor_by_plate: dict) -> pd.DataFrame:
    gs = group_summary.copy()

    if not exploratory:
        # rigorous mode: failed groups simply stay NaN downstream
        return gs

    for idx, row in gs[gs["QC_status"] == "FAILED"].iterrows():
        plate, sample, target = row["Plate"], row["Sample"], row["Target"]

        if fail_strategy == "drop":
            continue

        elif fail_strategy == "plate_mean":
            others = gs[(gs["Plate"] == plate) & (gs["Target"] == target) &
                        (gs["Sample"] != sample) & (gs["QC_status"] != "FAILED")]
            if others.empty or others["Cq_mean"].isna().all():
                LOG.warning("plate_mean imputation impossible for %s / %s / %s "
                            "(no passing samples on this plate for this target)",
                            plate, sample, target)
                continue
            val = float(others["Cq_mean"].mean())
            gs.loc[idx, "Cq_mean"] = val
            gs.loc[idx, "Imputed"] = True
            gs.loc[idx, "Imputed_note"] = (
                f"plate_mean: mean Cq of {len(others)} other passing "
                f"'{target}' samples on {plate} = {val:.3f}"
            )

        elif fail_strategy == "anchor_mean":
            anchor_sample = anchor_by_plate.get(plate)
            anchor_row = gs[(gs["Plate"] == plate) & (gs["Sample"] == anchor_sample) &
                             (gs["Target"] == target)]
            if anchor_row.empty or pd.isna(anchor_row["Cq_mean"].iloc[0]):
                LOG.warning("anchor_mean imputation impossible for %s / %s / %s "
                            "(anchor sample has no valid Cq for this target)",
                            plate, sample, target)
                continue
            val = float(anchor_row["Cq_mean"].iloc[0])
            gs.loc[idx, "Cq_mean"] = val
            gs.loc[idx, "Imputed"] = True
            gs.loc[idx, "Imputed_note"] = f"anchor_mean: Cq of anchor '{anchor_sample}' = {val:.3f}"

        elif fail_strategy == "best_subset":
            sub = raw[(raw["Plate"] == plate) & (raw["Sample"] == sample) &
                      (raw["Target"] == target)]
            cqs = sub["Cq"].dropna().tolist()
            wells = sub.loc[sub["Cq"].notna(), "Well Position"].tolist()
            if len(cqs) < 2:
                LOG.warning("best_subset imputation impossible for %s / %s / %s "
                            "(fewer than 2 valid Cq values at all)", plate, sample, target)
                continue
            # choose the 2-replicate pair with the smallest difference
            best_pair, best_sd = None, np.inf
            for i in range(len(cqs)):
                for j in range(i + 1, len(cqs)):
                    pair_sd = float(np.std([cqs[i], cqs[j]], ddof=1))
                    if pair_sd < best_sd:
                        best_sd, best_pair = pair_sd, (cqs[i], cqs[j])
            val = float(np.mean(best_pair))
            gs.loc[idx, "Cq_mean"] = val
            gs.loc[idx, "Cq_SD"] = best_sd
            gs.loc[idx, "Imputed"] = True
            gs.loc[idx, "Imputed_note"] = (
                f"best_subset: forced-pass on best 2 of {len(cqs)} replicates, "
                f"SD={best_sd:.3f} (still above threshold)"
            )
        else:
            raise ValueError(f"Unknown fail-strategy: {fail_strategy}")

    return gs


# 5. ddCt method

def compute_ddct(group_summary: pd.DataFrame, housekeeping: Optional[str],
                  anchor_by_plate: dict) -> pd.DataFrame:
    """ddCt vs. an anchor sample, per plate.

    If `housekeeping` is given, the standard two-step method is used:
        dCt  = Cq(target) - Cq(housekeeping)
        ddCt = dCt(sample) - dCt(anchor)

    If `housekeeping` is None, there is no reference-gene normalization step:
    each target's own Cq is used directly (dCt == Cq(target)) and ddCt is
    just Cq(target, sample) - Cq(target, anchor). This is the correct mode
    for experiments with no housekeeping/endogenous-control assay -- e.g.
    when template input was already equalized by another method (Qubit
    quantification, etc.) so no internal normalizer is needed or wanted.
    """
    rows = []

    for plate, plate_df in group_summary.groupby("Plate"):
        anchor_sample = anchor_by_plate.get(plate)
        if anchor_sample is None:
            raise ValueError(f"No anchor sample specified for plate '{plate}'.")

        if housekeeping is not None:
            hk = plate_df[plate_df["Target"] == housekeeping].set_index("Sample")["Cq_mean"]
            targets = sorted(t for t in plate_df["Target"].unique() if t != housekeeping)
        else:
            hk = None
            targets = sorted(plate_df["Target"].unique())

        anchor_dct = {}
        for target in targets:
            tgt_series = plate_df[plate_df["Target"] == target].set_index("Sample")["Cq_mean"]
            if anchor_sample not in tgt_series.index:
                anchor_dct[target] = np.nan
                continue
            if housekeeping is not None:
                if anchor_sample not in hk.index:
                    anchor_dct[target] = np.nan
                    continue
                anchor_dct[target] = tgt_series[anchor_sample] - hk[anchor_sample]
            else:
                anchor_dct[target] = tgt_series[anchor_sample]

        for target in targets:
            tgt_df = plate_df[plate_df["Target"] == target]
            hk_df = plate_df[plate_df["Target"] == housekeeping] if housekeeping is not None else None

            for _, row in tgt_df.iterrows():
                sample = row["Sample"]
                cq_tgt = row["Cq_mean"]

                if housekeeping is not None:
                    hk_row = hk_df[hk_df["Sample"] == sample]
                    cq_hk = hk_row["Cq_mean"].iloc[0] if not hk_row.empty else np.nan
                    dct = cq_tgt - cq_hk if pd.notna(cq_tgt) and pd.notna(cq_hk) else np.nan
                else:
                    cq_hk = np.nan
                    dct = cq_tgt if pd.notna(cq_tgt) else np.nan

                a_dct = anchor_dct.get(target, np.nan)
                ddct = dct - a_dct if pd.notna(dct) and pd.notna(a_dct) else np.nan
                rq = 2 ** (-ddct) if pd.notna(ddct) else np.nan
                log2fc = -ddct if pd.notna(ddct) else np.nan

                rows.append(dict(
                    Plate=plate, Sample=sample, Condition=row["Condition"],
                    BioRep=row["BioRep"], Target=target,
                    Anchor_sample=anchor_sample,
                    Housekeeping_gene=(housekeeping if housekeeping is not None
                                        else "none (direct Ct vs anchor)"),
                    Cq_target_mean=cq_tgt, Cq_target_SD=row["Cq_SD"],
                    Cq_housekeeping_mean=cq_hk,
                    QC_status_target=row["QC_status"], Imputed_target=row["Imputed"],
                    Imputed_note_target=row["Imputed_note"],
                    dCt=dct, anchor_dCt=a_dct, ddCt=ddct, RQ=rq, log2FC=log2fc,
                    is_anchor=(sample == anchor_sample),
                ))

    return pd.DataFrame(rows)


# 6. Statistics: normality -> parametric vs non-parametric -> vs control

@dataclass
class TargetStatsResult:
    target: str
    control_condition: str
    test_family: str                  
    omnibus_test: Optional[str] = None
    omnibus_p: Optional[float] = None
    normality: dict = field(default_factory=dict) 
    levene_p: Optional[float] = None
    posthoc: pd.DataFrame = None      
    note: Optional[str] = None


def run_stats_for_target(df_target: pd.DataFrame, control_condition: str,
                          alpha: float) -> TargetStatsResult:
    target_name = df_target["Target"].iloc[0]
    groups = {
        cond: sub["log2FC"].dropna().to_numpy()
        for cond, sub in df_target.groupby("Condition")
    }
    groups = {c: v for c, v in groups.items() if len(v) >= 1}

    if control_condition not in groups:
        return TargetStatsResult(target=target_name, control_condition=control_condition,
                                  test_family="insufficient_data",
                                  note="control condition has no usable log2FC values")

    if len(groups) < 2:
        return TargetStatsResult(target=target_name, control_condition=control_condition,
                                  test_family="insufficient_data",
                                  note="fewer than 2 conditions present for this target")

    # Conditions with <2 replicates can be compared with a
    # rank-based test, but variance-based methods (ANOVA/Dunnett) are not
    # valid for them -> force non-parametric 
    low_n = sorted(c for c, v in groups.items() if len(v) < 2)

    normality = {}
    all_normal = True
    for cond, vals in groups.items():
        if len(vals) >= 3:
            stat, p = sstats.shapiro(vals)
            normality[cond] = (float(stat), float(p))
            if p < alpha:
                all_normal = False
        else:
            normality[cond] = (np.nan, np.nan)
            all_normal = False  # can't confirm normality with n<3 -> be conservative

    groups_ge2 = {c: v for c, v in groups.items() if len(v) >= 2}
    levene_p = None
    if len(groups_ge2) >= 2:
        try:
            _, levene_p = sstats.levene(*groups_ge2.values())
        except Exception:
            levene_p = np.nan

    parametric = (all_normal and not low_n and
                  (levene_p is None or np.isnan(levene_p) or levene_p >= alpha))

    note_parts = []
    if low_n:
        note_parts.append(
            "condition(s) with <2 replicates (no within-group variance estimate; "
            f"forced non-parametric, inference for these is weak): {', '.join(low_n)}"
        )

    result = TargetStatsResult(
        target=target_name, control_condition=control_condition,
        test_family="parametric" if parametric else "nonparametric",
        normality=normality, levene_p=levene_p,
        note="; ".join(note_parts) or None,
    )

    other_conditions = [c for c in groups if c != control_condition]
    if not other_conditions:
        result.test_family = "insufficient_data"
        result.note = "no non-control conditions to compare"
        return result

    if parametric:
        try:
            f_stat, p_omni = sstats.f_oneway(*groups.values())
            result.omnibus_test = "one-way ANOVA"
            result.omnibus_p = float(p_omni)
        except Exception as exc:
            result.note = (result.note + " | " if result.note else "") + f"ANOVA failed: {exc}"

        try:
            dun = sstats.dunnett(
                *[groups[c] for c in other_conditions],
                control=groups[control_condition],
            )
            posthoc_rows = [
                dict(Condition=c, p_adj=float(p), stars=stars_for(p))
                for c, p in zip(other_conditions, dun.pvalue)
            ]
            result.posthoc = pd.DataFrame(posthoc_rows)
        except Exception as exc:
            result.note = (result.note + " | " if result.note else "") + f"Dunnett failed: {exc}"

    else:
        try:
            h_stat, p_omni = sstats.kruskal(*groups.values())
            result.omnibus_test = "Kruskal-Wallis"
            result.omnibus_p = float(p_omni)
        except Exception as exc:
            result.note = (result.note + " | " if result.note else "") + f"Kruskal-Wallis failed: {exc}"

        if sp is None:
            result.note = (result.note + " | " if result.note else "") + \
                "scikit-posthocs not installed; skipping Dunn's test"
        else:
            long_df = pd.concat(
                [pd.DataFrame({"Condition": c, "log2FC": v}) for c, v in groups.items()],
                ignore_index=True,
            )
            try:
                dunn = sp.posthoc_dunn(long_df, val_col="log2FC", group_col="Condition",
                                        p_adjust="holm")
                posthoc_rows = [
                    dict(Condition=c, p_adj=float(dunn.loc[control_condition, c]),
                         stars=stars_for(dunn.loc[control_condition, c]))
                    for c in other_conditions
                ]
                result.posthoc = pd.DataFrame(posthoc_rows)
            except Exception as exc:
                result.note = (result.note + " | " if result.note else "") + f"Dunn's test failed: {exc}"

    return result


def run_all_stats(ddct_df: pd.DataFrame, control_condition: str, alpha: float):
    results = {}
    for target, sub in ddct_df.groupby("Target"):
        results[target] = run_stats_for_target(sub, control_condition, alpha)
    return results

# 7. Plots

def _well_sort_key(well) -> tuple:
    """Sort key for a well position like 'A1', 'B12', ... by row letter then
    numeric column, so 'A2' sorts before 'A10' (unlike plain string sort)."""
    m = re.match(r"^([A-Za-z]+)0*(\d+)$", str(well).strip())
    if not m:
        return (str(well), 0)
    row, col = m.groups()
    return (row, int(col))


def compute_plate_sample_order(raw: pd.DataFrame) -> dict:
    """Map each plate -> samples ordered by their first well position on the
    physical plate (row letter then column number), i.e. the order the
    samples actually appear on the plate rather than an alphabetical or
    condition-based order. Used so plots read left-to-right the same way the
    plate is laid out, which makes visual QC/interpretation easier."""
    tmp = raw.copy()
    tmp["_well_key"] = tmp["Well Position"].map(_well_sort_key)
    order = {}
    for plate, pdf in tmp.groupby("Plate", sort=False):
        first_key = pdf.groupby("Sample")["_well_key"].min()
        order[plate] = first_key.sort_values().index.tolist()
    return order


def _sample_order(df: pd.DataFrame) -> list:
    """Fallback ordering (by already-resolved Condition, BioRep) for when no
    plate-layout order is available. Prefer compute_plate_sample_order()."""
    order_df = df[["Sample", "Condition", "BioRep"]].drop_duplicates()
    order_df = order_df.sort_values(["Condition", "BioRep", "Sample"])
    return order_df["Sample"].tolist()


def _fit_title_and_layout(fig, title_artist, ax=None, tight_kwargs: Optional[dict] = None,
                           margin_in: float = 0.5, max_iter: int = 4) -> None:
    """Widen `fig` in place, re-running tight_layout, until `title_artist`
    (from fig.suptitle(...) or ax.set_title(...)) fits within its reference
    width -- otherwise long plate/sample/target names get clipped or run off
    the edge of the saved image.

    A fig.suptitle(...) is centered on the whole FIGURE, so pass ax=None to
    compare its width against the figure width. An ax.set_title(...) is
    centered on that AXES (which is narrower than the figure once axis
    labels/ticks take up their own margin), so pass that ax -- comparing
    against the figure width alone under-widens and the title still clips.

    Call this once after setting the title, in place of a bare
    fig.tight_layout() call; it runs tight_layout itself (repeatedly, if
    needed) so the caller only needs to fig.savefig() afterwards."""
    tight_kwargs = tight_kwargs or {}
    if title_artist is None:
        fig.tight_layout(**tight_kwargs)
        return
    fig.tight_layout(**tight_kwargs)
    for _ in range(max_iter):
        try:
            fig.canvas.draw()
            renderer = fig.canvas.get_renderer()
            title_w_in = title_artist.get_window_extent(renderer=renderer).width / fig.dpi
            ref_w_in = (ax.get_window_extent(renderer=renderer).width / fig.dpi
                        if ax is not None else fig.get_size_inches()[0])
        except Exception:
            return
        deficit_in = title_w_in - ref_w_in
        if deficit_in <= 0:
            return
        cur_w, cur_h = fig.get_size_inches()
        fig.set_size_inches(cur_w + deficit_in + margin_in, cur_h, forward=True)
        fig.tight_layout(**tight_kwargs)


def plot_qc_replicates(replicate_log: pd.DataFrame, outdir: Path, dpi: int,
                        plate_sample_order: Optional[dict] = None):
    outdir.mkdir(parents=True, exist_ok=True)
    for plate, pdf in replicate_log.groupby("Plate"):
        targets = sorted(pdf["Target"].unique())
        samples = (plate_sample_order or {}).get(plate) or _sample_order(pdf)
        fig, axes = plt.subplots(len(targets), 1, figsize=(max(8, 0.5 * len(samples)), 4 * len(targets)),
                                  sharex=True)
        if len(targets) == 1:
            axes = [axes]

        for ax, target in zip(axes, targets):
            tdf = pdf[pdf["Target"] == target]
            for sample in samples:
                sdf = tdf[tdf["Sample"] == sample].sort_values("Well")
                x = [sample] * len(sdf)
                kept = sdf[~sdf["Removed"]]
                removed = sdf[sdf["Removed"] & sdf["Cq"].notna()]
                ax.scatter([sample] * len(kept), kept["Cq"], color="steelblue", zorder=3)
                ax.scatter([sample] * len(removed), removed["Cq"], color="crimson",
                           marker="x", s=90, zorder=4, label="Removed (outlier)")
            ax.set_ylabel(f"{target}\nCq")
            ax.grid(True, alpha=0.3)

        handles, labels = axes[0].get_legend_handles_labels()
        plt.setp(axes[-1].get_xticklabels(), rotation=60, ha="right")
        sup = fig.suptitle(f"QC - {plate}", y=0.995)
        if handles:
            fig.legend(handles[:1], labels[:1], loc="lower center",
                       bbox_to_anchor=(0.5, 0.0), ncol=1)
        _fit_title_and_layout(fig, sup, ax=None, tight_kwargs=dict(rect=[0, 0.05, 1, 0.96]))
        fig.savefig(outdir / f"qc_replicates_{plate}.png", dpi=dpi)
        plt.close(fig)


def plot_cq_by_sample(group_summary: pd.DataFrame, anchor_by_plate: dict,
                       outdir: Path, dpi: int,
                       plate_sample_order: Optional[dict] = None):
    outdir.mkdir(parents=True, exist_ok=True)
    for plate, pdf in group_summary.groupby("Plate"):
        anchor_sample = anchor_by_plate.get(plate)
        targets = sorted(pdf["Target"].unique())
        samples = (plate_sample_order or {}).get(plate) or _sample_order(pdf)

        variants = [("raw", False)]
        if anchor_sample is not None:
            variants.append(("with_anchor", True))

        for suffix, with_anchor in variants:
            fig, axes = plt.subplots(len(targets), 1,
                                      figsize=(max(8, 0.5 * len(samples)), 4 * len(targets)),
                                      sharex=True)
            if len(targets) == 1:
                axes = [axes]
            for ax, target in zip(axes, targets):
                tdf = pdf[pdf["Target"] == target].set_index("Sample").reindex(samples)
                colors = ["crimson" if (with_anchor and s == anchor_sample) else "steelblue"
                          for s in samples]
                ax.errorbar(samples, tdf["Cq_mean"], yerr=tdf["Cq_SD"], fmt="o",
                             ecolor="gray", capsize=3, color="steelblue", zorder=3)
                ax.scatter(samples, tdf["Cq_mean"], c=colors, zorder=4)
                if with_anchor and anchor_sample in tdf.index and pd.notna(tdf.loc[anchor_sample, "Cq_mean"]):
                    ax.axhline(tdf.loc[anchor_sample, "Cq_mean"], color="crimson",
                                linestyle="--", alpha=0.6,
                                label=f"anchor ({anchor_sample})")
                    ax.legend(fontsize=9)
                ax.set_ylabel(f"{target}\nCq (mean \u00b1 SD)")
                ax.grid(True, alpha=0.3)
            plt.setp(axes[-1].get_xticklabels(), rotation=60, ha="right")
            title = f"{plate} - Cq" + (" (anchor)" if with_anchor else "")
            sup = fig.suptitle(title)
            _fit_title_and_layout(fig, sup, ax=None, tight_kwargs=dict(rect=[0, 0, 1, 0.97]))
            fig.savefig(outdir / f"cq_by_sample_{suffix}_{plate}.png", dpi=dpi)
            plt.close(fig)


def _draw_significance_brackets(ax, x_positions: dict, control: str,
                                 posthoc: pd.DataFrame, y_top: float, y_step: float):
    if posthoc is None or posthoc.empty:
        return
    level = 0
    for _, row in posthoc.sort_values("Condition").iterrows():
        cond = row["Condition"]
        if cond not in x_positions or control not in x_positions:
            continue
        x1, x2 = x_positions[control], x_positions[cond]
        y = y_top + level * y_step
        ax.plot([x1, x1, x2, x2], [y, y + y_step * 0.15, y + y_step * 0.15, y],
                color="black", linewidth=1.2)
        ax.text((x1 + x2) / 2, y + y_step * 0.18, row["stars"],
                ha="center", va="bottom", fontsize=13)
        level += 1


def _draw_log2fc_panel(ax, sub: pd.DataFrame, control_condition: str,
                        res: Optional["TargetStatsResult"], title: str) -> None:
    """Draw one target's log2FC stripplot (+ mean markers, significance
    brackets vs. control, test-family annotation) onto an existing Axes.
    Shared by plot_log2fc (one file per target) and plot_log2fc_combined
    (all targets side by side in a single figure) so the two stay visually
    identical."""
    order = sorted(sub["Condition"].unique(),
                    key=lambda c: (c != control_condition, c))
    x_positions = {c: i for i, c in enumerate(order)}

    sns.stripplot(data=sub, x="Condition", y="log2FC", order=order,
                   ax=ax, size=8, jitter=0.15, alpha=0.85)
    means = sub.groupby("Condition")["log2FC"].mean().reindex(order)
    ax.scatter(range(len(order)), means, color="black", marker="_", s=800, zorder=5)
    ax.axhline(0, color="gray", linestyle=":", linewidth=1)
    ax.set_xticks(range(len(order)))
    ax.set_xticklabels(order, rotation=60, ha="right")
    ax.set_ylabel("log2FC")
    ax.set_title(title)

    if res is not None and res.posthoc is not None and not res.posthoc.empty:
        data_min, data_max = sub["log2FC"].min(), sub["log2FC"].max()
        data_range = max(data_max - data_min, 0.5)
        y_top = data_max + 0.15 * data_range
        y_step = 0.14 * data_range
        n_brackets = len(res.posthoc)
        _draw_significance_brackets(ax, x_positions, control_condition,
                                     res.posthoc, y_top, y_step)
        ax.set_ylim(data_min - 0.15 * data_range,
                    y_top + (n_brackets + 0.8) * y_step)

        fam = res.test_family
        omni = f"{res.omnibus_test}: p={res.omnibus_p:.4g}" if res.omnibus_p is not None else ""
        ax.text(0.01, 0.02, f"{fam}\n{omni}", transform=ax.transAxes,
                 ha="left", va="bottom", fontsize=9, color="dimgray")


def plot_log2fc(ddct_df: pd.DataFrame, stats_results: dict, control_condition: str,
                 outdir: Path, dpi: int):
    """One log2FC plot per target (primer pair), each its own file."""
    outdir.mkdir(parents=True, exist_ok=True)
    for target, sub in ddct_df.groupby("Target"):
        sub = sub.dropna(subset=["log2FC"])
        if sub.empty:
            continue
        n_conditions = sub["Condition"].nunique()
        fig, ax = plt.subplots(figsize=(max(6, 1.3 * n_conditions), 6))
        _draw_log2fc_panel(ax, sub, control_condition, stats_results.get(target),
                            f"{target} vs {control_condition}")
        _fit_title_and_layout(fig, ax.title, ax=ax)
        safe_target = re.sub(r"[^\w\-.]", "_", target)
        fig.savefig(outdir / f"log2FC_{safe_target}.png", dpi=dpi)
        plt.close(fig)


def plot_log2fc_combined(ddct_df: pd.DataFrame, stats_results: dict, control_condition: str,
                          outdir: Path, dpi: int):
    """All targets' log2FC panels side by side in a single figure, for an
    at-a-glance comparison across primer pairs -- in addition to (not
    instead of) the one-file-per-target plots from plot_log2fc()."""
    outdir.mkdir(parents=True, exist_ok=True)
    targets = [t for t, sub in ddct_df.groupby("Target")
               if not sub["log2FC"].dropna().empty]
    if not targets:
        return

    n_conditions = ddct_df["Condition"].nunique()
    panel_w = max(5, 1.1 * n_conditions)
    fig, axes = plt.subplots(1, len(targets), figsize=(panel_w * len(targets), 6))
    if len(targets) == 1:
        axes = [axes]

    for ax, target in zip(axes, targets):
        sub = ddct_df[ddct_df["Target"] == target].dropna(subset=["log2FC"])
        _draw_log2fc_panel(ax, sub, control_condition, stats_results.get(target), target)

    sup = fig.suptitle(f"log2FC vs {control_condition} - all targets", y=0.98)
    _fit_title_and_layout(fig, sup, ax=None, tight_kwargs=dict(rect=[0, 0, 1, 0.90]))
    fig.savefig(outdir / "log2FC_all_targets.png", dpi=dpi)
    plt.close(fig)


def _draw_rq_bar_panel(ax, sub: pd.DataFrame, control_condition: str,
                        res: Optional["TargetStatsResult"], title: str,
                        highlight_target: Optional[str] = None,
                        highlight_color: str = "#3B6FB8", other_color: str = "#9E9E9E",
                        highlight_label: str = "targets this gene",
                        other_label: str = "does not target this gene") -> Optional[list]:
    """Draw one target's RQ bar chart (mean bar + SEM-of-ddCt whiskers +
    individual biological-replicate points + dashed RQ=1 line + significance
    brackets vs. control) onto an existing Axes -- the same visual language
    as plot_log2fc's stripplot, just bar-style on the linear RQ scale
    instead of a scatter on the log2FC scale. Shared by plot_rq_bar (one
    file per target) and plot_rq_bar_combined (all targets side by side).

    Bar height = 2^-mean(ddCt) (the geometric mean of the per-replicate RQ
    values). Whiskers are +/- 1 SEM of ddCt, converted onto the RQ scale
    (asymmetric, since it's a log-linear transform) -- matching a standard
    ddCt-based RQ bar plot convention. Individual points are each biological
    replicate's own RQ = 2^-ddCt, jittered slightly on x.

    If `highlight_target` is given, a condition's bar/points are colored
    `highlight_color` when `highlight_target` appears as a case-insensitive
    substring of the condition name, else `other_color` -- reproduces the
    "targets this gene" / "does not target this gene" legend convention used
    for dCas9/CRISPRoff-style gene-silencing experiments. Leave it None for
    assays (e.g. digestion-efficiency) where that distinction doesn't apply;
    every bar is then drawn in `other_color`."""
    order = sorted(sub["Condition"].unique(), key=lambda c: (c != control_condition, c))
    x_positions = {c: i for i, c in enumerate(order)}

    rng = np.random.default_rng(0)
    bar_heights, err_low, err_high, colors = [], [], [], []
    point_x, point_y = [], []

    for c in order:
        csub = sub[sub["Condition"] == c]
        ddct_vals = csub["ddCt"].dropna().to_numpy()
        n = len(ddct_vals)
        x0 = x_positions[c]

        if n == 0:
            bar_heights.append(np.nan)
            err_low.append(0.0)
            err_high.append(0.0)
        else:
            mean_ddct = float(np.mean(ddct_vals))
            bar_rq = 2 ** (-mean_ddct)
            bar_heights.append(bar_rq)
            if n >= 2:
                sem_ddct = float(np.std(ddct_vals, ddof=1) / np.sqrt(n))
                err_low.append(bar_rq - 2 ** (-(mean_ddct + sem_ddct)))
                err_high.append(2 ** (-(mean_ddct - sem_ddct)) - bar_rq)
            else:
                err_low.append(0.0)
                err_high.append(0.0)
            jitter = rng.uniform(-0.12, 0.12, size=n)
            point_x.extend(x0 + jitter)
            point_y.extend(2 ** (-ddct_vals))

        if highlight_target:
            colors.append(highlight_color if highlight_target.lower() in c.lower() else other_color)
        else:
            colors.append(other_color)

    xs = [x_positions[c] for c in order]
    ax.bar(xs, bar_heights, color=colors, edgecolor="black", linewidth=1.1,
           width=0.62, zorder=2)
    ax.errorbar(xs, bar_heights, yerr=[err_low, err_high], fmt="none",
                ecolor="black", elinewidth=1.3, capsize=4, zorder=3)
    ax.scatter(point_x, point_y, facecolors="white", edgecolors="black",
               linewidth=1.1, s=55, zorder=4)
    ax.axhline(1.0, color="gray", linestyle="--", linewidth=1, zorder=1)

    ax.set_xticks(xs)
    ax.set_xticklabels(order, rotation=60, ha="right")
    ax.set_ylabel("RQ (2^−ΔΔCt)")
    ax.set_title(title)

    legend_handles = None
    if highlight_target:
        from matplotlib.patches import Patch
        legend_handles = [Patch(facecolor=other_color, edgecolor="black", label=other_label),
                           Patch(facecolor=highlight_color, edgecolor="black", label=highlight_label)]
        # Handed back to the caller rather than drawn here: the caller places
        # one shared legend above the figure (title row), clear of the
        # significance brackets/stats text that live inside the axes.

    valid_tops = [h + e for h, e in zip(bar_heights, err_high) if pd.notna(h)]
    visual_max = max(valid_tops + point_y + [1.0]) if (valid_tops or point_y) else 1.0
    valid_bottoms = [h - e for h, e in zip(bar_heights, err_low) if pd.notna(h)]
    visual_min = min([0.0] + valid_bottoms)

    if res is not None and res.posthoc is not None and not res.posthoc.empty:
        data_range = max(visual_max - visual_min, 0.3)
        y_top = visual_max + 0.15 * data_range
        y_step = 0.14 * data_range
        n_brackets = len(res.posthoc)
        _draw_significance_brackets(ax, x_positions, control_condition,
                                     res.posthoc, y_top, y_step)
        # Extra headroom (vs. the log2FC panel's +0.8) so the topmost
        # bracket's label doesn't collide with the test-family/omnibus-p
        # annotation anchored near the very top of the axes below.
        ax.set_ylim(min(0.0, visual_min - 0.1 * data_range),
                    y_top + (n_brackets + 1.9) * y_step)

        fam = res.test_family
        omni = f"{res.omnibus_test}: p={res.omnibus_p:.4g}" if res.omnibus_p is not None else ""
        ax.text(0.99, 0.995, f"{fam}\n{omni}", transform=ax.transAxes,
                 ha="right", va="top", fontsize=8, color="dimgray")
    else:
        ax.set_ylim(min(0.0, visual_min - 0.1 * max(visual_max - visual_min, 0.3)),
                    visual_max + 0.15 * max(visual_max - visual_min, 0.3))

    return legend_handles


_RQ_FOOTNOTE = ("RQ = relative quantification (2^−ΔΔCt), geometric mean of "
                "biological replicates; whiskers ± 1 SEM of ΔΔCt. "
                "Dashed line = {control} (RQ = 1).")


def plot_rq_bar(ddct_df: pd.DataFrame, stats_results: dict, control_condition: str,
                 outdir: Path, dpi: int, highlight_target_match: bool = False):
    """One RQ bar plot per target (primer pair) -- mean bar + SEM-of-ddCt
    whiskers + individual biological-replicate points, dashed RQ=1 line,
    significance brackets vs. control. Bar-chart counterpart to
    plot_log2fc()'s stripplot, on the linear RQ scale."""
    outdir.mkdir(parents=True, exist_ok=True)
    for target, sub in ddct_df.groupby("Target"):
        sub = sub.dropna(subset=["ddCt"])
        if sub.empty:
            continue
        n_conditions = sub["Condition"].nunique()
        res = stats_results.get(target)
        n_brackets = (len(res.posthoc) if (res is not None and res.posthoc is not None)
                      else max(0, n_conditions - 1))
        # More vertical room as the bracket stack grows -- a fixed height
        # squeezes the brackets/stars together until they're unreadable.
        fig_h = 6.5 + 0.6 * max(0, n_brackets - 2)
        fig, ax = plt.subplots(figsize=(max(6, 1.3 * n_conditions), fig_h))
        legend_handles = _draw_rq_bar_panel(
            ax, sub, control_condition, stats_results.get(target),
            f"{target} vs {control_condition}",
            highlight_target=(target if highlight_target_match else None))
        fig.text(0.01, 0.01, _RQ_FOOTNOTE.format(control=control_condition),
                  fontsize=7.5, color="dimgray", ha="left")
        top_rect = 1.0
        if legend_handles:
            fig.legend(handles=legend_handles, loc="upper center", ncol=2,
                       bbox_to_anchor=(0.5, 1.0), fontsize=9, frameon=False)
            top_rect = 0.90
        _fit_title_and_layout(fig, ax.title, ax=ax,
                               tight_kwargs=dict(rect=[0, 0.045, 1, top_rect]))
        safe_target = re.sub(r"[^\w\-.]", "_", target)
        fig.savefig(outdir / f"RQ_{safe_target}.png", dpi=dpi)
        plt.close(fig)


def plot_rq_bar_combined(ddct_df: pd.DataFrame, stats_results: dict, control_condition: str,
                          outdir: Path, dpi: int, highlight_target_match: bool = False):
    """All targets' RQ bar panels side by side in a single figure, mirroring
    plot_log2fc_combined() but on the linear RQ scale."""
    outdir.mkdir(parents=True, exist_ok=True)
    targets = [t for t, sub in ddct_df.groupby("Target")
               if not sub["ddCt"].dropna().empty]
    if not targets:
        return

    n_conditions = ddct_df["Condition"].nunique()
    # Wider than plot_log2fc_combined's panels: the stats annotation
    # (top-right) and the widest significance bracket (control vs. the
    # farthest condition, roughly centered) both compete for space at the
    # top of a bar-chart panel, more so than on the narrower stripplots.
    panel_w = max(6, 1.5 * n_conditions)
    max_brackets = 0
    for target in targets:
        res = stats_results.get(target)
        n_b = (len(res.posthoc) if (res is not None and res.posthoc is not None)
               else max(0, n_conditions - 1))
        max_brackets = max(max_brackets, n_b)
    # Same reasoning as plot_rq_bar: a fixed height squeezes the bracket
    # stack (and the significance stars on it) together until unreadable.
    fig_h = 6.5 + 0.6 * max(0, max_brackets - 2)
    fig, axes = plt.subplots(1, len(targets), figsize=(panel_w * len(targets), fig_h))
    if len(targets) == 1:
        axes = [axes]

    legend_handles = None
    for ax, target in zip(axes, targets):
        sub = ddct_df[ddct_df["Target"] == target].dropna(subset=["ddCt"])
        h = _draw_rq_bar_panel(ax, sub, control_condition, stats_results.get(target), target,
                                highlight_target=(target if highlight_target_match else None))
        legend_handles = legend_handles or h

    sup = fig.suptitle(f"RQ vs {control_condition} - all targets", y=0.985)
    fig.text(0.01, 0.01, _RQ_FOOTNOTE.format(control=control_condition),
              fontsize=7.5, color="dimgray", ha="left")
    top_rect = 0.95
    if legend_handles:
        fig.legend(handles=legend_handles, loc="upper center", ncol=2,
                   bbox_to_anchor=(0.5, 0.925), fontsize=9, frameon=False)
        top_rect = 0.88
    _fit_title_and_layout(fig, sup, ax=None, tight_kwargs=dict(rect=[0, 0.05, 1, top_rect]))
    fig.savefig(outdir / "RQ_all_targets.png", dpi=dpi)
    plt.close(fig)


# 8. Report

def write_stats_report(stats_results: dict, outpath: Path):
    lines = []
    for target, res in stats_results.items():
        lines.append(f"## Target: {target}")
        lines.append(f"Control condition: {res.control_condition}")
        lines.append(f"Test family: {res.test_family}")
        if res.note:
            lines.append(f"Note: {res.note}")
        if res.normality:
            lines.append("Shapiro-Wilk normality (per condition, n>=3 required):")
            for cond, (stat, p) in res.normality.items():
                p_str = f"{p:.4g}" if not np.isnan(p) else "n/a (n<3)"
                lines.append(f"  - {cond}: p={p_str}")
        if res.levene_p is not None:
            lp = f"{res.levene_p:.4g}" if not (res.levene_p is None or np.isnan(res.levene_p)) else "n/a"
            lines.append(f"Levene's test for equal variances: p={lp}")
        if res.omnibus_test:
            lines.append(f"Omnibus test: {res.omnibus_test}, p={res.omnibus_p:.4g}")
        if res.posthoc is not None and not res.posthoc.empty:
            lines.append("Post-hoc vs. control:")
            for _, row in res.posthoc.iterrows():
                lines.append(f"  - {row['Condition']}: p_adj={row['p_adj']:.4g} ({row['stars']})")
        lines.append("")
    outpath.write_text("\n".join(lines), encoding="utf-8")

# 9. CLI parsing and main entry point

def build_arg_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="Automated qPCR ddCt analysis: QC/outlier removal, ddCt, "
                    "RQ/log2FC, statistics, and plots.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    p.add_argument("--input", nargs="+", required=True, type=Path,
                    help="One or more QuantStudio export .xlsx files, one per plate.")
    p.add_argument("--plate-labels", nargs="+", default=None,
                    help="Custom plate labels, matched by order to --input "
                         "(default: file stem).")
    p.add_argument("--anchor", nargs="+", default=None,
                    help="Anchor/calibrator SAMPLE name (exact 'Sample' text, e.g. "
                         "'Empty 1' or 'Church wt'), one per --input file in the same "
                         "order. Giving --anchor is what triggers ddCt/RQ/log2FC + "
                         "statistics: with --housekeeping also set, ddCt uses the "
                         "standard dCt(target)-dCt(housekeeping) method referenced to "
                         "the anchor; without --housekeeping, ddCt is computed directly "
                         "from each target's own Cq referenced to the anchor (no "
                         "reference-gene normalization) -- appropriate when template "
                         "input was already equalized another way (e.g. Qubit-based "
                         "dilution), so there is no housekeeping/endogenous-control "
                         "assay on the plate. If --anchor is omitted entirely, no ddCt "
                         "math is done at all (QC + raw Cq outputs only).")
    p.add_argument("--control-condition", default=None,
                    help="Baseline condition group used for statistical comparisons "
                         "(default: condition resolved from the first plate's anchor "
                         "sample, e.g. 'E 1' -> 'E').")
    p.add_argument("--condition-regex", default=DEFAULT_CONDITION_REGEX,
                    help="Regex used to split a Sample name into biological condition "
                         "+ replicate id when no --sample-map entry or DA3 Biogroup is "
                         "available. Must contain a named group (?P<condition>...); a "
                         "(?P<rep>...) group is optional (defaults to '1'). "
                         f"Default: {DEFAULT_CONDITION_REGEX!r}")
    p.add_argument("--sample-map", type=Path, default=None,
                    help="Optional CSV with columns Sample,Condition,BioRep (and "
                         "optionally Plate) to explicitly declare which biological "
                         "condition/replicate each Sample belongs to. Takes priority "
                         "over DA3's own Biogroup field and over --condition-regex. "
                         "Use this when Sample names don't encode the replicate at all "
                         "and the DA3 'Biogroup' field was not set either.")
    p.add_argument("--housekeeping", default=None,
                    help="Housekeeping gene / Target name for the standard two-step "
                         "ddCt method (e.g. GAPDH). Omit this for experiments with no "
                         "relative-quantification reference gene: if --anchor is still "
                         "given, ddCt is computed directly against the anchor's own Cq "
                         "per target (no housekeeping subtraction) -- the right mode "
                         "when input was already equalized some other way (e.g. Qubit "
                         "dilution), so a fake reference-gene normalization would just "
                         "add noise. If --anchor is also omitted, the pipeline runs QC "
                         "+ raw Cq plots only and skips ddCt/RQ/log2FC and statistics "
                         "entirely.")
    p.add_argument("--sd-threshold", type=float, default=0.3,
                    help="Max allowed technical-replicate Cq SD before outlier "
                         "removal kicks in (default: 0.3). Ignored if "
                         "--skip-outlier-removal is set.")
    p.add_argument("--skip-outlier-removal", action="store_true", default=False,
                    help="Bypass the SD-based outlier removal step entirely: every "
                         "valid (non-NaN) technical replicate is kept and averaged "
                         "as-is, and --sd-threshold is ignored. A (sample, target) "
                         "group only FAILS if none of its replicates amplified at "
                         "all. Everything downstream (ddCt/RQ/log2FC, stats, plots) "
                         "still runs exactly as before -- use this when you just want "
                         "the raw-mean ddCt result without SD-based QC.")
    p.add_argument("--exploratory", action="store_true", default=False, # turned off by default
                    help="Enable imputation of FAILED (sample, target) groups "
                         "instead of dropping them. Off by default (rigorous mode).")
    p.add_argument("--fail-strategy", choices=["drop", "plate_mean", "anchor_mean", "best_subset"],
                    default="plate_mean",
                    help="How to salvage FAILED groups when --exploratory is set "
                         "(default: plate_mean). Ignored if --exploratory is not set.")
    p.add_argument("--alpha", type=float, default=0.05,
                    help="Significance threshold for normality/variance/omnibus tests "
                         "(default: 0.05).")
    p.add_argument("--rq-highlight-target-match", action="store_true", default=False,
                    help="In the RQ bar plots (plots/rq/), color a condition's bar/points "
                         "differently when the plot's target/gene name appears as a "
                         "case-insensitive substring of the condition name -- e.g. colors "
                         "'CRISPRoff-HER2' differently from 'CRISPRoff-RGS9' on the HER2 RQ "
                         "plot, reproducing the 'targets this gene' / 'does not target this "
                         "gene' legend convention used for dCas9/CRISPRoff-style "
                         "gene-silencing experiments. Off by default -- leave off for assays "
                         "(e.g. digestion-efficiency) where condition names don't encode a "
                         "target gene, since every bar would just end up 'does not target'.")
    p.add_argument("--dpi", type=int, default=150, help="Plot resolution (default: 150).")
    p.add_argument("--outdir", required=True, type=Path,
                    help="Output folder (created if it doesn't exist).")
    return p


def main(argv=None):
    args = build_arg_parser().parse_args(argv)

    if args.housekeeping is not None and not args.anchor:
        sys.exit("--anchor is required when --housekeeping is set "
                  "(the ddCt method needs a calibrator sample per plate).")

    if args.anchor is not None and len(args.anchor) != len(args.input):
        sys.exit(f"--anchor must have exactly one value per --input file "
                  f"({len(args.input)} inputs, {len(args.anchor)} anchors given).")

    plate_labels = args.plate_labels or [p.stem for p in args.input]
    if len(plate_labels) != len(args.input):
        sys.exit("--plate-labels must have exactly one value per --input file.")

    args.outdir.mkdir(parents=True, exist_ok=True)
    log_path = args.outdir / "run.log"
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(message)s",
        handlers=[logging.FileHandler(log_path, mode="w"), logging.StreamHandler(sys.stdout)],
    )
    LOG.info("qpcr_pipeline starting")
    LOG.info("Inputs: %s", [str(p) for p in args.input])
    LOG.info("Plate labels: %s", plate_labels)
    LOG.info("Anchors: %s", args.anchor)
    LOG.info("SD threshold: %s | skip_outlier_removal=%s | exploratory=%s | fail_strategy=%s",
              args.sd_threshold, args.skip_outlier_removal, args.exploratory, args.fail_strategy)

    anchor_by_plate = dict(zip(plate_labels, args.anchor)) if args.anchor else {}

    sample_map = load_sample_map(args.sample_map) if args.sample_map else None
    if sample_map is not None:
        LOG.info("Loaded --sample-map with %d entries from %s", len(sample_map), args.sample_map)

    # 1. read & combine
    raw_frames = []
    biogroup_maps = {}
    for path, label in zip(args.input, plate_labels):
        LOG.info("Reading %s as plate '%s'", path, label)
        raw_frames.append(read_plate_results(path, label))
        bg_map = read_biogroup_map(path)
        biogroup_maps[label] = bg_map
        if bg_map:
            LOG.info("Plate '%s': found DA3 Biogroup assignments for %d samples "
                      "-> will use these over --condition-regex", label, len(bg_map))
        else:
            LOG.info("Plate '%s': DA3 Biogroup field is empty/unset; falling back to "
                      "--sample-map (if given) then --condition-regex", label)
    raw = pd.concat(raw_frames, ignore_index=True)

    resolver = ConditionResolver(sample_map, biogroup_maps, args.condition_regex)

    # 2. QC / outlier resolution
    if args.skip_outlier_removal:
        LOG.info("--skip-outlier-removal set: averaging all valid replicates per "
                  "(sample, target) with no SD-based removal (--sd-threshold ignored)")
    else:
        LOG.info("Running SD-based outlier resolution (threshold=%s)", args.sd_threshold)
    replicate_log, group_summary = run_qc(raw, args.sd_threshold, resolver,
                                           skip_outlier_removal=args.skip_outlier_removal)
    LOG.info("Condition/replicate resolution sources used: %s", sorted(resolver.used_sources))

    n_failed = (group_summary["QC_status"] == "FAILED").sum()
    LOG.info("QC complete: %d/%d (sample, target) groups FAILED",
              n_failed, len(group_summary))

    # 3. exploratory imputation 
    group_summary = apply_fail_strategy(group_summary, raw, args.exploratory,
                                         args.fail_strategy, anchor_by_plate)
    if args.exploratory:
        n_imputed = group_summary["Imputed"].sum()
        LOG.info("Exploratory mode ('%s'): %d groups imputed", args.fail_strategy, n_imputed)

    # 4. ddCt (skipped entirely if --anchor was not given; --housekeeping just
    #    switches between the two-step and direct-vs-anchor ddCt method)
    ddct_df = None
    stats_results = {}
    control_condition = None

    if args.anchor:
        if args.housekeeping is not None:
            LOG.info("Computing ddCt / RQ / log2FC (housekeeping=%s, referenced to anchor)",
                      args.housekeeping)
        else:
            LOG.info("Computing ddCt / RQ / log2FC directly vs. anchor (no --housekeeping "
                      "given -- dCt = Cq(target) itself, no reference-gene normalization)")
        ddct_df = compute_ddct(group_summary, args.housekeeping, anchor_by_plate)

        control_condition = args.control_condition
        if control_condition is None:
            first_plate, first_anchor = plate_labels[0], args.anchor[0]
            control_condition, _, source = resolver.resolve(first_plate, first_anchor)
            LOG.info("--control-condition not given; resolved '%s' from first anchor "
                      "'%s' via %s", control_condition, first_anchor, source)

        # 5. statistics
        LOG.info("Running statistics vs. control condition '%s'", control_condition)
        stats_results = run_all_stats(ddct_df, control_condition, args.alpha)
    else:
        LOG.info("No --anchor given: skipping ddCt/RQ/log2FC and statistics -- "
                  "QC and raw Cq outputs only.")

    # 6. write tables 
    tables_dir = args.outdir / "tables"
    tables_dir.mkdir(parents=True, exist_ok=True)
    replicate_log.to_csv(tables_dir / "technical_replicates_QC.csv", index=False)
    group_summary.to_csv(tables_dir / "sample_target_Cq_summary.csv", index=False)

    if ddct_df is not None:
        ddct_df.to_csv(tables_dir / "ddCt_RQ_log2FC_results.csv", index=False)

        posthoc_frames = []
        for target, res in stats_results.items():
            if res.posthoc is not None and not res.posthoc.empty:
                f = res.posthoc.copy()
                f.insert(0, "Target", target)
                f.insert(1, "Control_condition", res.control_condition)
                f.insert(2, "Test_family", res.test_family)
                f.insert(3, "Omnibus_test", res.omnibus_test)
                f.insert(4, "Omnibus_p", res.omnibus_p)
                posthoc_frames.append(f)
        if posthoc_frames:
            pd.concat(posthoc_frames, ignore_index=True).to_csv(
                tables_dir / "statistics_posthoc.csv", index=False)
        write_stats_report(stats_results, tables_dir / "statistics_report.txt")

    # 7. plots
    LOG.info("Generating plots")
    plate_sample_order = compute_plate_sample_order(raw)
    plot_qc_replicates(replicate_log, args.outdir / "plots" / "qc", args.dpi,
                        plate_sample_order=plate_sample_order)
    plot_cq_by_sample(group_summary, anchor_by_plate, args.outdir / "plots" / "cq", args.dpi,
                       plate_sample_order=plate_sample_order)
    if ddct_df is not None:
        plot_log2fc(ddct_df, stats_results, control_condition, args.outdir / "plots" / "log2fc", args.dpi)
        plot_log2fc_combined(ddct_df, stats_results, control_condition,
                              args.outdir / "plots" / "log2fc", args.dpi)
        plot_rq_bar(ddct_df, stats_results, control_condition, args.outdir / "plots" / "rq", args.dpi,
                    highlight_target_match=args.rq_highlight_target_match)
        plot_rq_bar_combined(ddct_df, stats_results, control_condition, args.outdir / "plots" / "rq",
                              args.dpi, highlight_target_match=args.rq_highlight_target_match)

    LOG.info("Done. Results written to %s", args.outdir.resolve())
    print(f"\nDone. See:\n  {tables_dir}\n  {args.outdir / 'plots'}\n  {log_path}")


if __name__ == "__main__":
    main()